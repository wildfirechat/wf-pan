#!/usr/bin/env python3
"""生成「新建文档」用的空白模板：src/main/resources/doc-templates/new.{docx,xlsx,pptx}

不直接拿 ONLYOFFICE 自带的 new.* 模板（出自 ONLYOFFICE/document-templates，AGPL），
改用 python-docx / openpyxl / python-pptx（均为 MIT）生成，许可说明见同目录 NOTICE。
依赖：pip install python-docx==1.2.0 openpyxl==3.1.5 python-pptx==1.0.2
"""
import os
import sys

from docx import Document
from docx.oxml.ns import qn
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Emu

out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    os.path.dirname(__file__), '..', 'src', 'main', 'resources', 'doc-templates')
os.makedirs(out, exist_ok=True)

# 文字文档：正文中文字体指向随包分发的开源字体（Noto Sans SC），其余保持默认
doc = Document()
normal = doc.styles['Normal']
normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans SC')
doc.core_properties.author = ''
doc.core_properties.title = ''
doc.save(os.path.join(out, 'new.docx'))

wb = Workbook()
wb.active.title = 'Sheet1'
wb.properties.creator = ''
wb.save(os.path.join(out, 'new.xlsx'))

# 演示文稿：16:9，一页标题页
prs = Presentation()
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(6858000)
prs.slides.add_slide(prs.slide_layouts[0])
prs.core_properties.author = ''
prs.core_properties.title = ''
prs.save(os.path.join(out, 'new.pptx'))

print('ok', out)
